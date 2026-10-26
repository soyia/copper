#!/usr/bin/env python3
"""Build the M0 China trial paid-service document pack.

Design preset: compact_reference_guide.
Header pattern: customer_pack.
Named overrides: Chinese system font (Heiti SC), compact 10 pt table text,
and form-specific shaded entry cells. Page geometry remains US Letter with one-inch
margins and 9360 DXA usable width as required by the selected preset.
"""

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs/self-hosting/templates/Copper_中国境内首次收费与交付单据模板.docx"

PAGE_WIDTH_DXA = 12240
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 90
CELL_MARGIN_LEFT_RIGHT = 120

FONT_NAME = "Heiti SC"

NAVY = "1F3A5F"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHTER_BLUE = "F4F7FA"
LIGHT_GRAY = "F2F4F7"
ENTRY_FILL = "FAFAFA"
CAUTION_FILL = "FFF4CE"
WHITE = "FFFFFF"
INK = "222222"
MUTED = "666666"
GRID = "AEB7C2"


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_LEFT_RIGHT),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("end", CELL_MARGIN_LEFT_RIGHT),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Twips(width_dxa)


def set_table_geometry(table, widths: Sequence[int]):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        if row_pr.find(qn("w:cantSplit")) is None:
            row_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def format_paragraph(paragraph, after=3, before=0, line=1.15, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep


def clear_cell(cell):
    p = cell.paragraphs[0]
    p.clear()
    return p


def write_cell(cell, text, *, bold=False, size=10, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    p = clear_cell(cell)
    p.alignment = align
    format_paragraph(p, after=0, line=1.08)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        set_cell_shading(cell, fill)
    return cell


def add_table(doc, widths: Sequence[int], rows: int = 1):
    table = doc.add_table(rows=rows, cols=len(widths))
    set_table_geometry(table, widths)
    set_table_borders(table)
    table.style = "Table Grid"
    return table


def add_key_value_rows(doc, rows: Iterable[Sequence[str]], widths=(1900, 2780, 1900, 2780)):
    table = add_table(doc, widths, rows=0)
    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            label = idx % 2 == 0
            write_cell(
                row.cells[idx],
                value,
                bold=label,
                fill=LIGHT_BLUE if label else ENTRY_FILL,
                color=NAVY if label else INK,
            )
    set_table_geometry(table, widths)
    return table


def add_section_band(doc, title, code=None):
    table = add_table(doc, (7200, 2160), rows=1)
    write_cell(table.cell(0, 0), title, bold=True, size=11, color=WHITE, fill=NAVY)
    write_cell(table.cell(0, 1), code or "", bold=True, size=9, color=WHITE, fill=NAVY, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p = doc.add_paragraph()
    format_paragraph(p, after=1)
    return table


def add_note(doc, title, body, caution=False):
    table = add_table(doc, (1600, 7760), rows=1)
    write_cell(table.cell(0, 0), title, bold=True, color=NAVY, fill=CAUTION_FILL if caution else LIGHT_BLUE)
    write_cell(table.cell(0, 1), body, size=9.5, fill=CAUTION_FILL if caution else LIGHTER_BLUE)
    p = doc.add_paragraph()
    format_paragraph(p, after=2)
    return table


def add_checkbox_table(doc, headers, rows, widths):
    table = add_table(doc, widths, rows=1)
    for idx, header in enumerate(headers):
        write_cell(table.cell(0, idx), header, bold=True, color=NAVY, fill=LIGHT_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_row(table.rows[0])
    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            write_cell(row.cells[idx], value, size=9.5, fill=ENTRY_FILL if idx == len(values) - 1 else None)
    set_table_geometry(table, widths)
    return table


def add_signature_block(doc, left="服务提供方（盖章）", right="客户（盖章/签字）", compact=False):
    p = doc.add_paragraph()
    format_paragraph(p, after=0 if compact else 3)
    table = add_table(doc, (4680, 4680), rows=2 if compact else 3)
    write_cell(table.cell(0, 0), left, bold=True, fill=LIGHT_BLUE)
    write_cell(table.cell(0, 1), right, bold=True, fill=LIGHT_BLUE)
    if compact:
        write_cell(table.cell(1, 0), "签署人：________  日期：____年__月__日", fill=ENTRY_FILL)
        write_cell(table.cell(1, 1), "签署人：________  日期：____年__月__日", fill=ENTRY_FILL)
    else:
        write_cell(table.cell(1, 0), "签署人：\n\n", fill=ENTRY_FILL)
        write_cell(table.cell(1, 1), "签署人：\n\n", fill=ENTRY_FILL)
        write_cell(table.cell(2, 0), "日期：____年__月__日", fill=ENTRY_FILL)
        write_cell(table.cell(2, 1), "日期：____年__月__日", fill=ENTRY_FILL)
    return table


def add_form_title(doc, number, title, purpose):
    p = doc.add_paragraph()
    format_paragraph(p, after=2, keep=True)
    p.paragraph_format.page_break_before = True
    r = p.add_run(f"表 {number}")
    set_run_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    format_paragraph(p, after=5, keep=True)
    r = p.add_run(title)
    set_run_font(r, size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=10, keep=True)
    r = p.add_run(purpose)
    set_run_font(r, size=10.5, color=MUTED)


def add_dedicated_signature_page(doc, left, right):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    format_paragraph(p, after=5, keep=True)
    r = p.add_run("签署确认")
    set_run_font(r, size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=10, keep=True)
    r = p.add_run("本页为前述单据的组成部分。双方签署或盖章后分别留存，不得与正文分离使用。")
    set_run_font(r, size=10.5, color=MUTED)
    add_signature_block(doc, left, right)


def add_field_code(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_paragraph(p, after=0)
    r = p.add_run("COPPER  |  中国境内首次收费与交付单据")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(p, after=0)
    r = p.add_run("CN-TRIAL-20260713  |  第 ")
    set_run_font(r, size=9, color=MUTED)
    add_field_code(p, "PAGE")
    r = p.add_run(" 页 / 共 ")
    set_run_font(r, size=9, color=MUTED)
    add_field_code(p, "NUMPAGES")
    r = p.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)


def build_cover(doc):
    p = doc.add_paragraph()
    format_paragraph(p, after=2)
    r = p.add_run("CUSTOMER DELIVERY PACK")
    set_run_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    format_paragraph(p, after=6)
    r = p.add_run("Copper 中国境内首次收费与交付单据")
    set_run_font(r, size=26, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=20)
    r = p.add_run("适用于中国大陆境内小微企业、个体工商户及其他经营主体的低风险 B2B 试运营")
    set_run_font(r, size=12, color=MUTED)

    add_key_value_rows(
        doc,
        [
            ("条款版本", "CN-TRIAL-20260713", "模板版本", "M0-FORMS-20260713"),
            ("服务提供方", "填写公司全称", "统一社会信用代码", "填写"),
            ("适用订单", "部署、迁移、维护、代运维、定制与接口", "价格依据", "M0 目标定价目录"),
        ],
    )
    p = doc.add_paragraph()
    format_paragraph(p, after=6)
    add_note(
        doc,
        "使用原则",
        "每次收费至少填写表1；发生部署交付时填写表2和表3；需要远程接触客户数据时填写表4；服务结束或临时权限撤销时填写表5。双方签字或盖章后分别留存。",
    )
    add_note(
        doc,
        "重要边界",
        "本模板仅适用于中国大陆境内低风险 B2B 试运营。跨境数据、消费者交易、金融/医疗等高监管行业、公共SaaS、OEM、公开源码再发行和高金额特殊合同不得直接套用。",
        caution=True,
    )
    add_note(
        doc,
        "公司信息",
        "你已确认拥有公司实体且经营范围包含软件开发。首次使用前仍须把公司法定名称、统一社会信用代码、收款账户、纳税人身份和电子税务局可开票项目填写完整。",
    )


def build_order_form(doc):
    add_form_title(doc, "1", "Copper 技术服务订单 / 确认单", "用于确认交易主体、服务范围、含税价格、付款、发票、交付和双方责任。")
    add_section_band(doc, "A. 订单与交易主体", "必填")
    add_key_value_rows(
        doc,
        [
            ("订单编号", "CP-__________", "签订日期", "____年__月__日"),
            ("服务提供方", "公司法定全称", "统一社会信用代码", "填写"),
            ("注册地址", "填写", "联系人/电话", "填写"),
            ("收款账户", "户名 / 开户行 / 账号", "纳税人身份", "□小规模  □一般纳税人"),
            ("客户名称", "客户法定名称", "统一社会信用代码", "填写"),
            ("客户地址", "填写", "联系人/电话", "填写"),
        ],
    )
    add_section_band(doc, "B. 服务、环境与价格", "含税价")
    add_checkbox_table(
        doc,
        ("选择", "服务项目", "标准价格", "数量", "本单含税金额", "约定说明"),
        [
            ("□", "Copper 自助部署", "¥0", "__", "¥________", "客户自行部署维护"),
            ("□", "标准客户云部署", "¥349/生产环境", "__", "¥________", "兼容 Linux 环境"),
            ("□", "本地电脑 / NAS部署", "¥649起/环境", "__", "¥________", "仅限兼容清单"),
            ("□", "标准数据迁移", "¥249起/次", "__", "¥________", "先验样本再批量"),
            ("□", "安心维护", "¥349/年/环境", "__", "¥________", "4次，每次≤60分钟"),
            ("□", "代运维", "¥99/月或¥995/年", "__", "¥________", "基础监控与月度检查"),
            ("□", "非标准/定制/接口", "¥150/小时或固定报价", "__", "¥________", "另附需求与验收标准"),
        ],
        (560, 2100, 1580, 600, 1500, 3020),
    )
    add_key_value_rows(
        doc,
        [
            ("订单含税总价", "人民币（大写）：________  小写：¥________", "付款安排", "预付__% / 验收__%"),
            ("发票需求", "□普票  □专票  项目：________", "第三方费用", "□无  □客户直付  □另列清单"),
            ("生产环境", "数量__；地区__；系统/版本__", "计划交付", "____年__月__日"),
        ],
    )
    add_section_band(doc, "C. 交付内容与明确不包含事项", "须具体")
    add_checkbox_table(
        doc,
        ("交付内容", "约定结果/数量", "明确不包含或客户负责"),
        [
            ("安装与配置", "访问地址、版本、管理员初始化", "服务器采购、客户网络和非兼容系统改造"),
            ("HTTPS与域名", "域名：________", "域名/商业证书费用，除非订单另列"),
            ("备份与恢复", "目标：________；一次恢复验证", "长期巡检，除非购买维护/代运维"),
            ("数据迁移", "模板/数据范围：________", "未约定的清洗、历史单据和会计口径修正"),
            ("培训与文档", "对象/时长：________", "长期驻场、无限次数培训"),
            ("定制/接口", "附件编号：________", "未写入附件的需求、第三方接口费用"),
        ],
        (2200, 3400, 3760),
    )
    add_section_band(doc, "D. 关键条款确认", "显著提示")
    add_checkbox_table(
        doc,
        ("确认", "条款"),
        [
            ("□", "全部对外价格为人民币含税总价；除新增服务外不追加税点。"),
            ("□", "软件本身¥0；收费对应部署、迁移、维护、代运维、定制和接口等人工技术服务。"),
            ("□", "停止购买维护或代运维不锁定现有系统，不影响客户导出、备份和恢复自身数据。"),
            ("□", "客户持有服务器/云账号、数据库管理员凭据、备份和加密密钥；高风险操作需单独确认。"),
            ("□", "源码仅授权客户永久内部运行、修改、备份及委托保密服务商维护；未经授权不得转售、公共SaaS、OEM或使用Copper商标。"),
            ("□", "服务范围、验收期限、退款/重做、第三方费用和责任限制已经阅读并获得解释。"),
        ],
        (900, 8460),
    )
    add_note(doc, "适用条款", "本订单与《Copper 中国境内试运营条款》CN-TRIAL-20260713共同构成双方约定；本单特殊约定与通用条款冲突时，以双方明确签署的特殊约定为准。", caution=True)
    add_signature_block(doc)


def build_acceptance_form(doc):
    add_form_title(doc, "2", "Copper 部署交付与验收单", "用于证明实际部署、备份恢复验证、管理员交接、遗留风险和客户验收结果。")
    add_section_band(doc, "A. 基本信息", "关联订单")
    add_key_value_rows(
        doc,
        [
            ("订单编号", "CP-__________", "验收单编号", "AC-__________"),
            ("客户名称", "填写", "生产环境", "□客户云  □本地  □NAS"),
            ("Copper版本", "填写版本/提交号", "交付日期", "____年__月__日"),
            ("访问地址", "填写", "服务器区域", "填写"),
        ],
    )
    add_section_band(doc, "B. 交付检查", "逐项验证")
    add_checkbox_table(
        doc,
        ("结果", "检查项", "证据/说明"),
        [
            ("□通过 □不适用", "服务可启动，健康检查通过，服务器重启后可恢复运行", "时间/截图编号：________"),
            ("□通过 □不适用", "管理员首次初始化完成，无公共默认密码", "管理员账号：仅记录名称，不记录密码"),
            ("□通过 □不适用", "Web、API、数据库、附件和反向代理拓扑与订单一致", "拓扑/端口：________"),
            ("□通过 □不适用", "数据库、附件、配置和备份目录已固定并说明", "路径清单：________"),
            ("□通过 □不适用", "HTTPS、域名和证书状态符合约定", "有效期/续期主体：________"),
            ("□通过 □不适用", "关键业务流程抽样：登录、建档、查询、保存和导出", "抽样结果：________"),
            ("□通过 □不适用", "日志、磁盘和备份失败查看方式已说明", "命令/位置：________"),
        ],
        (1500, 4300, 3560),
    )
    add_section_band(doc, "C. 备份与恢复实证", "未验证不得完整交付")
    add_key_value_rows(
        doc,
        [
            ("备份生成时间", "____年__月__日 __:__", "备份位置", "填写客户控制路径"),
            ("备份范围", "□数据库  □附件  □配置  □清单", "校验值/大小", "填写"),
            ("恢复目标", "□空白环境  □隔离测试库", "恢复完成时间", "____年__月__日 __:__"),
            ("恢复验证", "□可登录  □数据可查  □附件可读", "关键数量/余额", "填写抽样结果"),
            ("最终结论", "□通过  □有条件通过  □未通过", "证据位置", "截图/日志/报告编号"),
        ],
    )
    add_section_band(doc, "D. 遗留问题与验收结论", "双方确认")
    add_checkbox_table(
        doc,
        ("编号", "遗留问题/已知风险", "责任人", "计划完成", "状态"),
        [
            ("1", "", "", "", "□待办 □接受风险"),
            ("2", "", "", "", "□待办 □接受风险"),
            ("3", "", "", "", "□待办 □接受风险"),
        ],
        (600, 4100, 1400, 1700, 1560),
    )
    add_key_value_rows(
        doc,
        [
            ("验收结果", "□通过  □有条件通过  □不通过", "整改期限", "____年__月__日"),
            ("客户意见", "填写；可另附页", "复验日期", "____年__月__日"),
        ],
    )
    add_signature_block(doc, "交付人/服务提供方（盖章）", "客户验收人（盖章/签字）")


def build_handover_form(doc):
    add_form_title(doc, "3", "Copper 管理员账号、数据与密钥交接单", "记录控制权已经移交客户。不得在本单填写密码、私钥、恢复码或完整密钥值。")
    add_note(doc, "安全警告", "本单只记录账号名称、资产归属、安全交付渠道和客户验证结果。密码、私钥、恢复码和备份加密密钥必须通过双方认可的安全渠道单独传递。", caution=True)
    add_key_value_rows(doc, [("订单编号", "CP-__________", "客户名称", "填写"), ("交接日期", "____年__月__日", "交接执行人", "填写")])
    add_checkbox_table(
        doc,
        ("资产", "标识/位置（不含秘密值）", "归属", "安全交付方式", "客户验证"),
        [
            ("服务器/云主账号", "账号名/控制台地址", "客户", "□现场 □密码管理器 □其他", "□已登录"),
            ("域名与DNS", "域名/服务商", "客户", "□现场 □安全链接 □其他", "□可管理"),
            ("数据库管理员", "账号名/主机/端口", "客户", "□现场 □密码管理器 □其他", "□已连接"),
            ("Copper管理员", "管理员账号名", "客户", "□首次设置 □安全链接", "□已登录"),
            ("备份目录/仓库", "路径/存储桶", "客户", "□权限授予 □控制台", "□可读取"),
            ("备份加密密钥", "密钥编号/保管位置", "客户", "□密码管理器 □离线介质", "□已验证"),
            ("证书与自动续期", "证书域名/续期主体", "客户", "□账号移交 □自动化", "□已确认"),
            ("发布包与源码", "版本/校验值/位置", "客户许可", "□仓库 □加密介质", "□可读取"),
        ],
        (1600, 2400, 950, 2500, 1910),
    )
    add_section_band(doc, "临时权限与撤销计划", "最小权限")
    add_checkbox_table(
        doc,
        ("临时账号/权限", "用途", "到期时间", "撤销责任人", "撤销验证"),
        [("", "", "", "", "□已撤销"), ("", "", "", "", "□已撤销"), ("", "", "", "", "□已撤销")],
        (2200, 2200, 1700, 1800, 1460),
    )
    add_dedicated_signature_page(doc, "交付人/服务提供方（盖章）", "接收人/客户（盖章/签字）")


def build_data_authorization_form(doc):
    add_form_title(doc, "4", "Copper 数据处理授权与远程服务记录", "客户授权远程部署、迁移、排障或代运维时使用；一次服务一份记录。")
    add_key_value_rows(
        doc,
        [
            ("授权编号", "DA-__________", "关联订单/工单", "填写"),
            ("客户名称", "填写", "服务执行人", "填写"),
            ("授权开始", "____年__月__日 __:__", "最晚结束", "____年__月__日 __:__"),
            ("处理目的", "□部署 □迁移 □排障 □代运维 □其他", "访问方式", "□远程桌面 □SSH □控制台"),
        ],
    )
    add_section_band(doc, "授权范围", "必要最小")
    add_checkbox_table(
        doc,
        ("选择", "数据/权限类型", "必要范围与限制"),
        [
            ("□", "系统配置与运行日志", "仅查看与故障相关内容，日志脱敏后再导出"),
            ("□", "客户/供应商/联系人", "仅处理约定问题，不用于画像、营销或模型训练"),
            ("□", "商品、订单、库存", "只处理工单列明范围，不改变已确认业务结果"),
            ("□", "应收应付、资金信息", "高敏数据；修改、导出或覆盖必须再次书面确认"),
            ("□", "附件与备份", "只复制必要样本；约定保存期限后返还或删除"),
            ("□", "数据库/服务器管理", "使用临时最小权限；危险操作单独确认"),
        ],
        (700, 2900, 5760),
    )
    add_note(doc, "默认限制", "未经客户书面同意不得转委托、不得传输到境外、不得用于训练模型或无关分析、不得长期保留客户数据；服务结束后应撤销权限并返还或删除临时副本。", caution=True)
    add_section_band(doc, "操作记录与结束处理", "可追溯")
    add_checkbox_table(
        doc,
        ("时间", "执行人", "操作/命令/数据范围", "结果与证据"),
        [("", "", "", ""), ("", "", "", ""), ("", "", "", ""), ("", "", "", "")],
        (1400, 1400, 3800, 2760),
    )
    add_key_value_rows(
        doc,
        [
            ("临时副本", "□未产生  □已返还  □已删除", "删除/返还时间", "____年__月__日 __:__"),
            ("临时权限", "□已撤销  □保留至____", "安全事件", "□无  □有，事件编号____"),
            ("服务结果", "□完成  □部分完成  □失败", "客户确认", "□已复核操作记录"),
        ],
    )
    add_signature_block(doc, "服务执行人/提供方（签字）", "授权人/客户（签字或盖章）")


def build_closeout_form(doc):
    add_form_title(doc, "5", "Copper 服务完结与权限撤销确认单", "用于部署、迁移、排障、维护或代运维阶段结束时确认资料返还、权限撤销和剩余责任。")
    add_key_value_rows(
        doc,
        [
            ("完结单编号", "CL-__________", "订单/工单", "填写"),
            ("客户名称", "填写", "完结日期", "____年__月__日"),
            ("服务类型", "□部署 □迁移 □维护 □代运维 □定制", "完结原因", "□完成 □取消 □终止"),
        ],
    )
    add_section_band(doc, "完结检查", "逐项确认")
    add_checkbox_table(
        doc,
        ("结果", "检查项", "证据/说明"),
        [
            ("□完成 □不适用", "客户已取得服务器、数据库、Copper管理员和备份控制权", "交接单编号：________"),
            ("□完成 □不适用", "服务提供方临时账号、SSH密钥、远程桌面和控制台授权已撤销", "撤销时间：________"),
            ("□完成 □不适用", "临时数据、日志、备份和附件副本已返还或删除", "记录编号：________"),
            ("□完成 □不适用", "未结费用、第三方续费和发票事项已列明", "说明：________"),
            ("□完成 □不适用", "遗留问题、风险、责任人和计划日期已经交接", "清单编号：________"),
            ("□完成 □不适用", "停止维护不锁定系统；后续支持按新订单或既有有效维护期执行", "客户已知悉"),
        ],
        (1500, 5000, 2860),
    )
    add_section_band(doc, "遗留事项与最终结论", "不得留空")
    add_checkbox_table(
        doc,
        ("编号", "遗留事项/风险", "责任方", "计划日期", "处理方式"),
        [("1", "", "", "", ""), ("2", "", "", "", ""), ("3", "", "", "", "")],
        (600, 3900, 1500, 1600, 1760),
    )
    add_key_value_rows(
        doc,
        [
            ("最终结论", "□服务完结  □有遗留但同意完结  □未完结", "后续联系人", "填写"),
            ("客户备注", "填写；可另附页", "资料附件", "□有  □无"),
        ],
    )
    add_dedicated_signature_page(doc, "服务提供方（盖章）", "客户（盖章/签字）")


def build():
    doc = Document()
    configure_document(doc)
    build_cover(doc)
    build_order_form(doc)
    build_acceptance_form(doc)
    build_handover_form(doc)
    build_data_authorization_form(doc)
    build_closeout_form(doc)

    props = doc.core_properties
    props.title = "Copper 中国境内首次收费与交付单据模板"
    props.subject = "M0 中国境内低风险 B2B 试运营"
    props.author = "Copper"
    props.keywords = "Copper, 技术服务订单, 部署验收, 密钥交接, 数据处理授权"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
